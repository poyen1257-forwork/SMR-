from pathlib import Path
from tempfile import gettempdir

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont


OUTPUT = Path("docs/SMR系統架構程式說明文件.docx")

FONT_CJK = "Microsoft JhengHei"
FONT_LATIN = "Aptos"
COLOR_NAVY = "17365D"
COLOR_BLUE = "D9EAF7"
COLOR_PALE = "F4F8FC"
COLOR_GRID = "D9D9D9"
COLOR_TEXT = "1F1F1F"


def architecture_font(size, bold=False):
    font_name = "msjhbd.ttc" if bold else "msjh.ttc"
    font_path = Path("C:/Windows/Fonts") / font_name
    return ImageFont.truetype(str(font_path), size)


def draw_architecture_flow() -> Path:
    """Create a compact architecture diagram embedded in the Word document."""
    image = Image.new("RGB", (1800, 980), "white")
    draw = ImageDraw.Draw(image)
    title_font = architecture_font(38, bold=True)
    label_font = architecture_font(26, bold=True)
    body_font = architecture_font(22)
    small_font = architecture_font(20)

    draw.text((72, 46), "SMR 控制架構流程圖", fill="#000000", font=title_font)
    draw.text((72, 103), "感測資料由 Xsens 傳至上位機，再由 ESP32 透過 CAN 控制馬達", fill="#555555", font=body_font)

    def box(x, y, width, height, title, subtitle, fill="#F4F8FC"):
        radius = 18
        draw.rounded_rectangle(
            (x, y, x + width, y + height), radius=radius, fill=fill, outline="#17365D", width=3
        )
        title_box = draw.textbbox((0, 0), title, font=label_font)
        subtitle_box = draw.multiline_textbbox((0, 0), subtitle, font=body_font, spacing=5, align="center")
        title_x = x + (width - (title_box[2] - title_box[0])) / 2
        title_y = y + 25
        subtitle_x = x + (width - (subtitle_box[2] - subtitle_box[0])) / 2
        subtitle_y = y + 82
        draw.text((title_x, title_y), title, fill="#17365D", font=label_font)
        draw.multiline_text((subtitle_x, subtitle_y), subtitle, fill="#1F1F1F", font=body_font, spacing=5, align="center")

    def arrow(x1, y1, x2, y2, label=None):
        draw.line((x1, y1, x2, y2), fill="#3776A6", width=5)
        direction = 1 if x2 >= x1 else -1
        if abs(x2 - x1) >= abs(y2 - y1):
            points = [(x2, y2), (x2 - 18 * direction, y2 - 11), (x2 - 18 * direction, y2 + 11)]
        else:
            direction = 1 if y2 >= y1 else -1
            points = [(x2, y2), (x2 - 11, y2 - 18 * direction), (x2 + 11, y2 - 18 * direction)]
        draw.polygon(points, fill="#3776A6")
        if label:
            label_box = draw.textbbox((0, 0), label, font=small_font)
            label_x = (x1 + x2 - (label_box[2] - label_box[0])) / 2
            label_y = (y1 + y2 - (label_box[3] - label_box[1])) / 2 - 34
            draw.text((label_x, label_y), label, fill="#3776A6", font=small_font)

    top_y = 250
    box(70, top_y, 225, 150, "MTi-630R", "九軸姿態\n量測", "#E8F2FB")
    box(355, top_y, 255, 150, "mti630r_reader.py", "XDA 資料讀取\n100 Hz", "#E8F2FB")
    box(955, top_y, 255, 150, "xsens2esp32.py", "yaw 差值計算\n相對位置命令", "#FFF4E5")
    box(1270, top_y, 200, 150, "ESP32", "Arduino 韌體\nSerial 命令", "#FFF4E5")

    arrow(295, 325, 355, 325, "感測封包")
    arrow(610, 325, 955, 325)
    json_note = (
        "TCP/IP 傳輸，port 5006\n"
        'JSON：{"time_s": ..., "imu_tx_ns": ...,\n'
        '       "roll": ..., "pitch": ..., "yaw": ...}'
    )
    draw.multiline_text((650, 420), json_note, fill="#3776A6", font=small_font, spacing=4)
    arrow(1210, 325, 1270, 325, "USB Serial")

    bottom_y = 610
    box(1045, bottom_y, 230, 150, "WCMCU230", "CAN 收發器\nCANH / CANL", "#FFF4E5")
    box(1335, bottom_y, 250, 150, "MF4015v2", "位置與速度\n馬達控制", "#F2F8EE")
    arrow(1370, 400, 1160, 610, "GPIO43 / GPIO44")
    arrow(1275, 685, 1335, 685, "CAN 1 Mbps\nID 0x141")

    box(175, 600, 560, 165, "control_system.py", "start：啟動遠端 reader 與本機控制\nstatus：查看執行狀態    end：停止 bridge 與 reader", "#F6F0FA")
    draw.line((455, 600, 455, 470), fill="#7B5AA6", width=3)
    draw.line((455, 470, 485, 470), fill="#7B5AA6", width=3)
    draw.line((485, 470, 485, 400), fill="#7B5AA6", width=3)
    draw.polygon([(485, 400), (474, 418), (496, 418)], fill="#7B5AA6")
    draw.text((195, 792), "控制層只做 yaw 決策；ESP32 只做馬達 CAN 通訊與命令轉換。", fill="#555555", font=small_font)

    output = Path(gettempdir()) / "smr_architecture_flow.png"
    image.save(output)
    return output


def set_run_font(run, size=None, bold=None, color=None, font_name=None):
    name = font_name or FONT_LATIN
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, color=COLOR_GRID, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_cell_width(cell, width_inches):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_inches * 1440)))
    tc_w.set(qn("w:type"), "dxa")


def set_table_layout_fixed(table):
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run("SMR 程式運作框架 | ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    paragraph._p.append(field)


def style_paragraph(paragraph, space_before=0, space_after=6, line_spacing=1.35):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(space_before)
    fmt.space_after = Pt(space_after)
    fmt.line_spacing = line_spacing


def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    style_paragraph(p, space_after=7)
    if bold_prefix and text.startswith(bold_prefix):
        run = p.add_run(bold_prefix)
        set_run_font(run, size=10.5, bold=True, color=COLOR_TEXT)
        run = p.add_run(text[len(bold_prefix):])
        set_run_font(run, size=10.5, color=COLOR_TEXT)
    else:
        run = p.add_run(text)
        set_run_font(run, size=10.5, color=COLOR_TEXT)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(16 if level == 1 else 10)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_run_font(run, size=15 if level == 1 else 12, bold=True, color="000000")
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.1
        run = p.add_run(item)
        set_run_font(run, size=10, color=COLOR_TEXT)


def add_numbered(doc, items):
    for index, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.65)
        p.paragraph_format.first_line_indent = Cm(-0.55)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.25
        run = p.add_run(f"{index}.  ")
        set_run_font(run, size=10.5, color=COLOR_TEXT)
        run = p.add_run(item)
        set_run_font(run, size=10.5, color=COLOR_TEXT)


def add_code(doc, code):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Cm(0.45)
    p.paragraph_format.right_indent = Cm(0.45)
    p.paragraph_format.line_spacing = 1.12
    p_pr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "F2F2F2")
    p_pr.append(shd)
    run = p.add_run(code)
    set_run_font(run, size=9, color="333333", font_name="Consolas")
    return p


def add_table(doc, headers, rows, widths, font_size=9.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    set_table_layout_fixed(table)
    header = table.rows[0]
    repeat_header(header)
    prevent_row_split(header)
    for index, value in enumerate(headers):
        cell = header.cells[index]
        set_cell_width(cell, widths[index])
        set_cell_shading(cell, COLOR_NAVY)
        set_cell_border(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style_paragraph(p, space_after=0, line_spacing=1.1)
        run = p.add_run(value)
        set_run_font(run, size=font_size, bold=True, color="FFFFFF")

    for row_index, values in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        for index, value in enumerate(values):
            cell = row.cells[index]
            set_cell_width(cell, widths[index])
            set_cell_border(cell)
            set_cell_margins(cell)
            if row_index % 2:
                set_cell_shading(cell, COLOR_PALE)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            style_paragraph(p, space_after=0, line_spacing=1.14)
            run = p.add_run(value)
            set_run_font(run, size=font_size, color=COLOR_TEXT)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def build_document():
    architecture_diagram = draw_architecture_flow()
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.7)
    section.left_margin = Cm(1.85)
    section.right_margin = Cm(1.85)
    section.header_distance = Cm(0.8)
    section.footer_distance = Cm(0.8)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = FONT_LATIN
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
    normal.font.size = Pt(10.5)
    for name in ("Title", "Heading 1", "Heading 2"):
        style = styles[name]
        style.font.name = FONT_LATIN
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CJK)
        style.font.color.rgb = RGBColor(0, 0, 0)

    footer = section.footer
    footer_p = footer.paragraphs[0]
    add_page_number(footer_p)
    for run in footer_p.runs:
        set_run_font(run, size=8.5, color="666666")

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(7)
    run = title.add_run("SMR 程式運作框架")
    set_run_font(run, size=23, bold=True, color="000000")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(20)
    run = subtitle.add_run("Xsens MTi 630R 姿態量測與 ESP32 CAN 馬達控制系統")
    set_run_font(run, size=11, color="555555")

    add_body(
        doc,
        "本文件說明 SMR 系統現行程式的職責分工、資料流、通訊格式、控制流程與操作方式。系統以 Xsens MTi-630R 提供 yaw 姿態資料，再經 Python 控制程式與 ESP32 CAN 韌體驅動 MF4015v2 馬達。",
    )

    add_heading(doc, "1. 系統目的與分層")
    add_body(
        doc,
        "SMR 的核心目標是依感測器姿態角度進行馬達角度追蹤。系統將感測資料讀取、網路傳輸、控制決策與 CAN 馬達命令分開處理。ESP32 僅處理馬達通訊與 CAN frame；感測器驅動、yaw 差值計算及速度決策均由 Python 執行。",
    )
    add_table(
        doc,
        ["層級", "元件或程式", "主要責任"],
        [
            ("感測層", "Xsens MTi-630R", "量測三軸加速度、角速度、磁場及姿態角；目前控制主要使用 yaw。"),
            ("資料讀取", "mti630r_reader.py", "透過 Xsens 官方 XDA 函式庫讀取感測資料，建立 TCP JSON 串流。"),
            ("控制邏輯", "xsens2esp32.py", "將 yaw 變化轉為相對位置命令，經 USB Serial 傳送至 ESP32。"),
            ("系統管理", "control_system.py", "管理遠端 reader 與本機 bridge 的 start、status、end 流程。"),
            ("馬達韌體", "mf5015v2_sample.ino", "接收 Serial 指令，產生 CAN frame，並讀取馬達狀態。"),
            ("通訊與致動", "WCMCU230 與 MF4015v2", "將 ESP32 邏輯訊號轉為 CANH/CANL，並由馬達驅動器執行控制。"),
        ],
        [1.05, 2.05, 3.9],
    )

    add_heading(doc, "2. 整體資料與控制流程")
    add_body(doc, "系統的主要資料流如下。控制資料由感測器向馬達傳遞；CAN 狀態回覆與延遲資訊則由馬達回傳至 Python。")
    figure = doc.add_paragraph()
    figure.alignment = WD_ALIGN_PARAGRAPH.CENTER
    figure.paragraph_format.space_after = Pt(4)
    figure.add_run().add_picture(str(architecture_diagram), width=Inches(6.75))
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(9)
    run = caption.add_run("圖 1  SMR 感測與馬達控制架構")
    set_run_font(run, size=9.5, color="555555")
    add_table(
        doc,
        ["順序", "來源", "傳遞資料", "目的端"],
        [
            ("1", "MTi-630R", "感測器封包", "mti630r_reader.py"),
            ("2", "mti630r_reader.py", "TCP JSON：roll、pitch、yaw、imu_tx_ns", "xsens2esp32.py"),
            ("3", "xsens2esp32.py", "Serial：n <角度> <速度>", "ESP32"),
            ("4", "ESP32", "CAN frame，1 Mbps，標準 ID 0x141", "MF4015v2"),
            ("5", "MF4015v2 與 ESP32", "CAN 回覆與 LAT 時間資料", "xsens2esp32.py"),
        ],
        [0.6, 1.55, 3.0, 1.85],
    )
    add_body(doc, "control_system.py 不計算 yaw，也不直接產生 CAN 命令；它只負責將 reader 與 xsens2esp32.py 啟動、監看及結束。")

    add_heading(doc, "3. Xsens 感測器資料讀取")
    add_body(
        doc,
        "mti630r_reader.py 透過 xsensdeviceapi 開啟 MTi-630R，從 callback 持續取得最新封包。程式可輸出完整九軸資料至 CSV，也可透過 TCP 或 UDP 對外送出姿態資料。提供控制程式的 TCP 封包為一筆一行的 JSON。",
    )
    add_code(doc, '{"time_s":1750000000.123,"imu_tx_ns":1750000000123456789,"roll":0.12,"pitch":-0.35,"yaw":45.67}')
    add_table(
        doc,
        ["欄位", "意義", "單位"],
        [
            ("time_s", "人類閱讀用的 Unix 秒時間。", "s"),
            ("imu_tx_ns", "reader 準備送出 TCP JSON 時記錄的 Unix 奈秒時間。", "ns"),
            ("roll、pitch、yaw", "姿態角。控制程式目前只使用 yaw。", "degree"),
        ],
        [1.35, 4.45, 1.2],
    )
    add_body(doc, "預設感測資料讀取與 TCP 串流頻率皆為 100 Hz。雖然控制目前只使用 yaw，封包仍保留 roll 與 pitch，方便後續擴充多軸控制或進行姿態資料檢查。")

    add_heading(doc, "4. xsens2esp32 控制邏輯")
    add_body(
        doc,
        "xsens2esp32.py 是 yaw 跟隨控制的核心。它預設連接 Xsens TCP 10.196.173.57:5006，並以 COM7、115200 baud 與 ESP32 通訊。未使用 --arm 時為 Dry Run，只顯示決策結果，不會開啟 COM port 或移動馬達。",
    )
    add_heading(doc, "4.1 控制步驟", level=2)
    add_numbered(
        doc,
        [
            "接收 TCP JSON，確認 yaw 是有效數值。",
            "以第一筆 yaw 建立基準角度；armed 模式會先送出 s，確保馬達停止。",
            "將 yaw 差值限制在 -180 至 +180 度，避免跨越正負 180 度時誤判為 360 度轉動。",
            "當累積角度差達 deadband，且已超過最短命令間隔時，建立相對位置命令。",
            "以角度差除以經過時間估算 yaw 速度，乘上 gain 後限制在馬達速度範圍內。",
            "將命令寫入 ESP32，並將本次 yaw 更新為新的控制基準。",
        ],
    )
    add_table(
        doc,
        ["參數", "預設值", "作用"],
        [
            ("--deadband", "0.05 degree", "抑制小於門檻的角度變化，減少感測器雜訊造成的頻繁命令。"),
            ("--command-period", "0.01 s", "最短命令間隔，最高約 100 Hz 更新。"),
            ("--gain", "1.2", "放大估算的 yaw 速度，使馬達較能跟上感測器動作。"),
            ("--min-speed", "10 dps", "位置控制的最小速度。"),
            ("--max-speed", "720 dps", "位置控制的速度上限。"),
            ("--speed-step", "1 dps", "對速度四捨五入，降低命令抖動。"),
        ],
        [1.65, 1.15, 4.2],
    )
    add_body(doc, "一般位置命令格式如下。角度為相對於軟體維護目標位置的變化量，而非每次直接指定馬達機械零點。")
    add_code(doc, "n <相對角度> <最大速度>\n範例：n -0.25 30")
    add_body(doc, "watchdog 預設為 0.5 秒。超過此時間未收到 yaw 時，程式會提示資料逾時，但目前仍保留最後一個位置目標，不會自動送出停止命令。若需要資料中斷立即停止，需由操作端送 s，或在後續版本加入 fail-safe 停止策略。")

    add_heading(doc, "5. ESP32 與 CAN 馬達控制")
    add_body(doc, "ESP32 韌體使用 Arduino Core 內建的 driver/twai.h，不需額外安裝 ESP32CAN 函式庫。TWAI 設定為 1 Mbps、標準 CAN ID 0x141。")
    add_table(
        doc,
        ["XIAO ESP32-S3", "WCMCU230 或馬達端", "用途"],
        [
            ("D6 / GPIO43", "CTX / TXD", "ESP32 TX 輸出至 CAN 收發器。"),
            ("D7 / GPIO44", "CRX / RXD", "CAN 收發器 RX 輸出至 ESP32。"),
            ("3V3", "VCC", "WCMCU230 邏輯電源。"),
            ("GND", "GND 與馬達電源負極", "共地參考。"),
            ("WCMCU230 CANH/CANL", "馬達 CANH/CANL", "CAN 差動通訊線。"),
        ],
        [1.7, 2.35, 2.95],
    )
    add_body(doc, "啟動時韌體先檢查 D6/D7 不被拉低，再設定 D6 為輸出、D7 為輸入上拉，並將 TX drive capability 設為最低等級，以降低 GPIO 輸出負擔。Arduino IDE 必須啟用 USB CDC On Boot。")
    add_heading(doc, "5.1 Serial 命令與 CAN 對應", level=2)
    add_table(
        doc,
        ["Serial 命令", "功能", "CAN 命令"],
        [
            ("r", "讀取溫度、匯流排電壓與錯誤碼。", "0x9A"),
            ("m <dps>", "連續速度控制，範圍 -720 至 720 dps。", "0xA2"),
            ("n <deg> <dps>", "相對位置控制，速度範圍 1 至 720 dps。", "先讀 0x92，再送 0xA4"),
            ("s", "停止馬達；速度模式漸減，位置模式或靜止時直接停止。", "0x81"),
        ],
        [1.35, 3.7, 1.95],
    )
    add_body(doc, "第一個 n 命令會先使用 0x92 讀取馬達多圈角度，作為 position_target_cdeg 初始值。後續相對角度會換算為 0.01 度單位，累積成絕對目標位置，再以 0xA4 送至馬達。")
    add_body(doc, "讀取、速度與停止命令會等待 TX ACK 及對應回覆。若出現 TX failed: no CAN ACK 或 TWAI state=2，代表 CAN 未收到可確認的節點或已進入 BUS_OFF，應檢查馬達供電、CANH/CANL、終端電阻、鮑率及收發器狀態，而不是持續重送位置命令。")

    add_heading(doc, "6. 延遲紀錄機制")
    add_body(doc, "每隔 10 筆位置命令，bridge 會送出含 sequence 與時間戳記的延伸 n 命令。ESP32 在接收序列命令及確認 CAN 成功送出時記錄時間，並以 LAT 回覆。結果寫入 logs/yaw_motor_latency.csv。")
    add_code(doc, "n <角度> <速度> <sequence> <imu_tx_ns> <bridge_tx_ns>\nLAT,<sequence>,<imu_tx_ns>,<bridge_tx_ns>,<esp_rx_us>,<esp_can_tx_us>")
    add_table(
        doc,
        ["CSV 欄位", "來源", "意義"],
        [
            ("imu_tx_ns", "Xsens reader", "JSON 輸出前的時間戳記。"),
            ("bridge_rx_ns", "xsens2esp32.py", "bridge 收到 JSON 的時間。"),
            ("bridge_tx_ns", "xsens2esp32.py", "bridge 寫入 Serial 前的時間。"),
            ("esp_rx_us", "ESP32", "ESP32 收到延伸命令時的單調計時。"),
            ("esp_can_tx_us", "ESP32", "確認 CAN frame 成功送出時的單調計時。"),
            ("imu_to_bridge_ms", "CSV 計算", "IMU JSON 送出至 bridge 收到的延遲。"),
            ("bridge_to_esp_ack_ms", "CSV 計算", "bridge 寫入 Serial 到收到 LAT 回覆的延遲。"),
            ("imu_to_esp_ack_ms", "CSV 計算", "IMU 時間戳到收到 ESP32 回覆的端到端延遲。"),
            ("esp_serial_to_can_tx_ms", "CSV 計算", "ESP32 收到序列命令到 CAN 成功送出的時間。"),
        ],
        [2.05, 1.55, 3.4],
        font_size=9,
    )
    add_body(doc, "imu_to_bridge_ms 與 imu_to_esp_ack_ms 需要兩台電腦時間同步；若時鐘不同步，應以 --imu-clock-offset-ms 修正。上述數值量到的是通訊與命令送出延遲，並非馬達機械軸實際開始轉動或到達目標的延遲；後者需以馬達編碼器回授與誤差門檻量測。")

    add_heading(doc, "7. control_system 系統管理")
    add_body(doc, "control_system.py 是一般操作入口。它透過 SSH 連接 Xsens 所在電腦，啟動或停止遠端 reader，並在本機啟動 xsens2esp32.py。SSH 密碼於執行時輸入或由 IMU_SSH_PASSWORD 環境變數提供，不寫入原始碼。")
    add_table(
        doc,
        ["輸入", "行為"],
        [
            ("start", "停止既有遠端 reader，建立新的 reader，確認 TCP 5006 監聽後啟動本機 xsens2esp32.py。"),
            ("status", "顯示本機 bridge 是否執行，以及遠端 reader 是否正在監聽。"),
            ("end", "要求 bridge 結束，嘗試送出最後的 ESP32 s，再停止遠端 reader。"),
            ("quit", "執行與 end 相同的結束流程後離開。"),
        ],
        [1.35, 5.65],
    )
    add_body(doc, "bridge 輸出會在控制器視窗以前綴 BRIDGE> 顯示，同時寫入 logs/xsens2esp32.log。")

    add_heading(doc, "8. 程式碼函式與參數詳解")
    add_body(doc, "本章以目前儲存庫的實作為準，說明函式的輸入、輸出與責任。Python 中的輸出值通常透過 return 回傳；ESP32 韌體則以 bool 表示是否成功，並將診斷訊息輸出到 USB Serial。所有角度均使用 degree，速度使用 dps（degree per second），CAN 位置資料以 0.01 degree 為最小單位。")

    add_heading(doc, "8.1 mti630r_reader.py 感測資料讀取", level=2)
    add_body(doc, "此程式是 Xsens 官方 XDA API 與控制系統間的介面。它讀取感測器封包、轉成 Python 字典，選擇性寫入 CSV，並建立 TCP server 讓 xsens2esp32.py 主動連線取得一行一筆的 JSON。")
    add_table(
        doc,
        ["命令列參數", "預設值", "輸入限制", "用途"],
        [
            ("--rate", "100 Hz", "1 至 400", "要求 MTi 輸出的姿態與九軸資料頻率。"),
            ("--yaw-only", "關閉", "旗標", "只要求姿態資料；啟用時不要求加速度、陀螺儀與磁場資料。"),
            ("--print-rate", "10 Hz", "> 0", "限制終端機列印頻率，不影響感測器本身的讀取頻率。"),
            ("--stream-rate", "100 Hz", "> 0", "非 yaw-only 模式下 TCP/UDP JSON 的最大傳送頻率。"),
            ("--duration", "0 s", ">= 0", "0 代表持續執行；正值代表執行指定秒數後結束。"),
            ("--csv", "未設定", "檔案路徑", "設定時，將完整九軸與姿態資料寫為 CSV。"),
            ("--udp-host / --udp-port", "未設定 / 5005", "合法 IP 與 1 至 65535", "設定後額外以 UDP 傳送 JSON。"),
            ("--tcp-port", "未設定", "1 至 65535", "建立 TCP server；本系統使用 5006。"),
        ],
        [1.45, 1.15, 1.55, 2.85],
        font_size=8.8,
    )
    add_table(
        doc,
        ["函式或方法", "輸入", "輸出", "用途與內部處理"],
        [
            ("PacketCallback.__init__", "max_packets: int，預設 20", "PacketCallback 物件", "建立固定長度 deque 與 Lock。deque 滿時會淘汰最舊封包，避免讀取端落後造成延遲持續累積。"),
            ("PacketCallback.onLiveDataAvailable", "XDA 裝置、XDA packet", "無", "XDA callback。將收到的封包複製為 XsDataPacket 後，在鎖定區塊中放入佇列。"),
            ("PacketCallback.get_packet", "無", "XsDataPacket 或 None", "以 thread-safe 方式取出最早等待處理的封包；佇列空時回傳 None。"),
            ("find_mti", "無", "XsPortInfo", "使用 XsScanner_scanPorts 掃描連接埠，回傳第一個 MTi 或 MTi-G；找不到時拋出 RuntimeError。"),
            ("configure_device", "device、rate_hz、yaw_only", "無；失敗時 RuntimeError", "切至設定模式，要求 PacketCounter、SampleTimeFine、Quaternion；非 yaw-only 時再要求加速度、角速度與磁場，最後切回量測模式。"),
            ("packet_values", "一筆 XDA packet", "dict 或 None", "讀取 Euler roll、pitch、yaw；若有 calibrated data，再寫入 acc、gyro、mag 三軸數值。方向資料不可用時回傳 None。"),
            ("parse_args", "命令列參數", "argparse.Namespace", "宣告 reader 的所有參數並檢查頻率、時間與通訊埠範圍。"),
            ("main", "解析後的命令列設定", "int：0 成功，1 失敗", "掃描與開啟 Xsens、建立 callback、依設定開啟 CSV/UDP/TCP，持續輸出 JSON，並在 finally 關閉 socket、port 與 XDA control。"),
        ],
        [1.42, 1.42, 1.3, 3.26],
        font_size=8.4,
    )
    add_body(doc, "輸出 JSON 由 stream_values 建立，欄位為 time_s、imu_tx_ns、roll、pitch、yaw。json.dumps(..., allow_nan=False) 會拒絕 NaN，避免將無效姿態送進控制端；TCP 傳送時在 JSON 後加上換行字元，讓接收端能辨識每一筆資料的結尾。")

    add_heading(doc, "8.2 xsens2esp32.py yaw 控制 bridge", level=2)
    add_body(doc, "此程式接收 reader 的 JSON，將相鄰 yaw 差轉為相對位置命令 n，再經 USB Serial 交給 ESP32。它不是直接以 yaw 絕對值控制馬達，因此可持續累積跨多圈的動作。")
    add_table(
        doc,
        ["命令列參數", "預設值", "輸入限制", "用途"],
        [
            ("--transport", "tcp", "tcp 或 udp", "選擇接收 yaw 的傳輸方式；目前系統使用 TCP。"),
            ("--sensor-ip / --sensor-port", "10.196.173.57 / 5006", "IP / 1 至 65535", "TCP 模式主動連線的 reader 位址。"),
            ("--listen-host / --listen-port", "0.0.0.0 / 5005", "IP / 1 至 65535", "僅 UDP 模式使用的本機監聽位址。"),
            ("--esp-port / --baud", "COM7 / 115200", "Serial port / 正整數", "ESP32 USB Serial 連線設定。"),
            ("--arm", "關閉", "旗標", "未啟用時為 Dry Run，不開啟 COM port，也不送馬達命令。"),
            ("--deadband", "0.05 degree", ">= 0", "yaw 累積變化達此門檻才送出位置命令，抑制雜訊。"),
            ("--gain", "1.2", "> 0", "將估算 yaw 速度乘上此值，讓馬達速度更能追上感測器。"),
            ("--min-speed / --max-speed", "10 / 720 dps", "1 至 720", "命令速度的下限與上限；實際送出的速度會被夾在此範圍。"),
            ("--command-period", "0.01 s", "> 0", "兩次位置命令最短間隔；0.01 s 對應最高約 100 Hz。"),
            ("--speed-step", "1 dps", "1 至 max-speed", "速度以此步進四捨五入，避免微小速度差造成 Serial 命令抖動。"),
            ("--watchdog", "0.5 s", "> 0", "超過此時間未收到 yaw 時輸出警告；現行程式保留最後位置目標，不會自動停馬達。"),
            ("--latency-sample-every", "10", ">= 1", "每 N 筆位置命令加入時間戳，要求 ESP32 回傳 LAT 資料。"),
            ("--latency-csv", "logs/yaw_motor_latency.csv", "檔案路徑", "延遲量測資料的 CSV 位置。"),
            ("--imu-clock-offset-ms", "0 ms", "浮點數", "控制端時鐘減感測端時鐘的校正值；兩台電腦未同步時才需調整。"),
            ("--duration", "0 s", ">= 0", "0 代表直到 Ctrl+C；正值代表指定執行秒數。"),
        ],
        [1.58, 1.28, 1.42, 2.57],
        font_size=8.2,
    )
    add_table(
        doc,
        ["函式", "輸入", "輸出", "用途與內部處理"],
        [
            ("wrap_degrees", "angle: float", "float，-180 至 <180", "以 modulo 將角度差折回最短方向；例如 179 到 -179 度會得到 +2 度，而不是 -358 度。"),
            ("parse_args", "命令列參數", "argparse.Namespace", "建立控制參數，並驗證 deadband、gain、速度、port、週期與延遲取樣設定。"),
            ("open_motor", "args", "serial.Serial 或 None", "只有 --arm 時才開啟 ESP32 COM port。DTR/RTS 固定為 False，避免 XIAO ESP32-S3 因控制線被重設或進入 boot 模式。"),
            ("send_motor", "port、command: str、echo", "無", "在 armed 模式將 ASCII 命令加換行寫入 Serial；echo 為真時同時印出 MOTOR> 訊息。"),
            ("read_yaw", "一筆 bytes JSON", "(yaw, imu_tx_ns) 或 None", "UTF-8 解碼並解析 JSON，取出有限值 yaw 與可選的 imu_tx_ns；格式或數值錯誤直接忽略。"),
            ("record_esp32_timing", "Serial port、pending dict、CSV writer、buffer、clock offset", "未處理完的 bytes buffer", "非阻塞讀取 ESP32 輸出。一般訊息印出；LAT 回覆則與 sequence 對應、計算延遲並寫入 CSV。"),
            ("main", "命令列設定", "int：0 成功，1 連線或開埠失敗", "開 TCP/UDP、建立第一筆 yaw 基準、計算 delta 與速度，送出 n 命令，Ctrl+C 或結束時送 s 並關閉資源。"),
        ],
        [1.4, 1.58, 1.42, 3.0],
        font_size=8.35,
    )
    add_body(doc, "main 內的核心計算為 delta = wrap_degrees(yaw - last_motion_yaw)，yaw_speed_dps = abs(delta) / elapsed，speed = clamp(round(yaw_speed_dps × gain)，min-speed，max-speed)。若 delta 未達 deadband 或距上次命令不足 command-period，程式不送新命令。一般命令格式為 n <delta_degrees> <speed>；延遲取樣命令再附加 sequence、imu_tx_ns、bridge_tx_ns。")

    add_heading(doc, "8.3 control_system.py 系統管理", level=2)
    add_body(doc, "此程式是互動式啟動器。它透過 SSH 在 Xsens 電腦執行 reader，再在本機以 subprocess 啟動 xsens2esp32.py。它管理流程，但不解析 IMU 封包，也不直接組 CAN frame。")
    add_table(
        doc,
        ["命令列參數", "預設值", "用途"],
        [
            ("--remote-host", "10.196.173.57", "Xsens reader 所在電腦的 SSH 與 TCP 位址。"),
            ("--remote-user", "pmc64", "遠端 Windows 登入帳號。"),
            ("--remote-root", "C:\\Users\\PMC64\\Desktop\\xsens_mti630r", "遠端 reader.py、log 檔所在資料夾。"),
            ("--remote-python", "Python39 python.exe", "遠端執行 mti630r_reader.py 的直譯器完整路徑。"),
            ("--sensor-port", "5006", "傳入 reader 與 bridge 的 TCP port。"),
            ("--esp-port", "COM7", "傳入 bridge 與最終停止命令使用的 ESP32 Serial port。"),
            ("--deadband / --max-speed", "0.05 / 720", "傳遞給 bridge 的控制安全限制。"),
            ("--latency-sample-every", "10", "傳遞給 bridge 的延遲取樣間隔。"),
            ("--arm", "關閉", "傳遞給 bridge；未啟用時系統只做 Dry Run。"),
        ],
        [1.75, 2.0, 3.65],
        font_size=8.8,
    )
    add_table(
        doc,
        ["方法", "輸入", "輸出", "用途與內部處理"],
        [
            ("YawMotorSystem.__init__", "args、SSH password", "系統物件", "保存設定、密碼、本機 bridge Process、log handle 與輸出轉送 thread。"),
            ("remote_powershell", "PowerShell script: str", "遠端 stdout: str", "以 UTF-16LE Base64 封裝指令後透過 Paramiko SSH 執行；SSH 連線 timeout 為 10 秒。"),
            ("stop_remote_reader", "無", "無", "尋找 command line 含 mti630r_reader.py 的遠端行程並停止，輸出停止數量。"),
            ("start_remote_reader", "無", "bool", "先停止舊 reader，再以 rate 100、stream-rate 100、tcp-port 啟動新行程；最多等待約 5 秒確認 TCP Listen。"),
            ("forward_bridge_output", "無，使用 self.bridge", "無", "背景 thread 持續讀 bridge stdout，寫入 logs/xsens2esp32.log 並在畫面前加 BRIDGE>。"),
            ("close_bridge_log", "無", "無", "等待輸出 thread 結束並關閉本機 log 檔。"),
            ("start_bridge", "無", "bool", "以目前 Python 直譯器啟動 python/xsens2esp32.py，帶入感測 IP、Serial port、deadband、max-speed 與 --arm。"),
            ("send_stop_to_esp", "無", "無", "短暫開啟 ESP32 Serial，送出 s\\n 後關閉；DTR/RTS 保持 False。"),
            ("stop_bridge", "無", "無", "先以 CTRL_BREAK_EVENT 請 bridge 正常退出，逾時後 terminate/kill，接著送 s 並關閉 log。"),
            ("start", "無", "無", "依序啟動遠端 reader 與本機 bridge，並捕捉 SSH 或網路例外。"),
            ("end", "無", "無", "停止 bridge、嘗試停止遠端 reader，最後輸出 System ended。"),
            ("status", "無", "無", "列出本機 bridge 是否執行，並查詢遠端 TCP port 是否 Listen。"),
            ("parse_args / main", "命令列與互動輸入", "int：0", "驗證系統設定，取得 SSH 密碼後讀取 start、status、end、quit 命令。"),
        ],
        [1.45, 1.45, 1.05, 3.45],
        font_size=8.15,
    )

    add_heading(doc, "8.4 mf5015v2_sample.ino 韌體", level=2)
    add_body(doc, "ESP32 韌體將簡短的 Serial 文字命令轉成 8-byte CAN frame。使用 ESP-IDF TWAI driver，在 XIAO ESP32-S3 上以 D6/GPIO43 為 TX、D7/GPIO44 為 RX，透過 WCMCU230 的 CTX/CRX 進入 CAN 匯流排。")
    add_table(
        doc,
        ["常數或狀態", "值或型別", "用途"],
        [
            ("CAN_TX_PIN / CAN_RX_PIN", "GPIO43 / GPIO44", "CAN 控制器連到 WCMCU230 的 CTX 與 CRX。"),
            ("MOTOR_CAN_ID", "0x141", "送往 Driver ID 0 的標準 CAN identifier。"),
            ("READ_MULTI_TURN_ANGLE / READ_STATE_1", "0x92 / 0x9A", "讀取絕對多圈位置，或讀取溫度、電壓、錯誤狀態。"),
            ("SPEED_CONTROL / POSITION_CONTROL_SPEED / MOTOR_STOP", "0xA2 / 0xA4 / 0x81", "連續速度、帶速度的絕對位置、停止命令。"),
            ("MIN_SPEED_DPS / MAX_SPEED_DPS", "-720 / 720", "m 命令可接受的連續速度範圍。"),
            ("MIN_POSITION_SPEED_DPS / MAX_POSITION_SPEED_DPS", "1 / 720", "n 命令可接受的位置控制最大速度範圍。"),
            ("can_ready", "bool", "start_can 成功後為 true；所有 CAN 操作先檢查此狀態。"),
            ("commanded_speed_dps / soft_stop_active", "int / bool", "保存最後速度模式速度與漸停是否進行中。"),
            ("position_target_initialized / position_target_cdeg", "bool / int32", "保存是否已讀取起始多圈角度，以及目前累積的絕對目標位置，單位 0.01 degree。"),
            ("command_buffer / command_length", "char[96] / size_t", "USB Serial 的行緩衝；長度超過 95 字元會拒絕該命令。"),
        ],
        [2.15, 1.75, 4.05],
        font_size=8.35,
    )
    add_table(
        doc,
        ["函式", "輸入", "輸出", "用途與內部處理"],
        [
            ("print_frame", "label、twai_message_t", "無", "以十六進位列印 CAN ID、DLC 與最多 8 bytes 資料，供接線與協定除錯。"),
            ("print_status", "無", "無", "讀取 TWAI status，列印 state、TX/RX error counter、bus error count。"),
            ("start_can", "無", "bool", "先以 INPUT_PULLUP 檢查 D6/D7 未被拉低；再設 D6 OUTPUT、D7 INPUT_PULLUP、最低 TX drive capability，安裝並啟動 1 Mbps normal-mode TWAI。"),
            ("can_exchange", "data[8]、expected_command、reply 指標", "bool", "清除舊 RX frame，送出 frame，最多等待 500 ms 的 TX ACK，再最多等待 1 s 的同 ID/命令回覆；成功時寫入 reply。"),
            ("can_send_position_target", "data[8]、wait_for_tx_success、can_tx_us 指標", "bool", "快速送出位置 frame。一般追蹤時不等待 ACK；延遲取樣時最多等待 30 ms，並記錄 CAN 成功送出的 ESP 微秒時間。"),
            ("read_motor_state", "無", "無", "以 0x9A 呼叫 can_exchange，解析 reply[1] 溫度、reply[2:4] 乘 0.01 V 的電壓與 reply[7] 錯誤碼。"),
            ("run_motor", "speed_dps: int", "bool", "將 dps 乘 100 成為協定的 0.01 dps/LSB，低位元優先寫入 0xA2 的 bytes 4 至 7。"),
            ("read_motor_position", "position_cdeg 指標", "bool", "以 0x92 讀取 reply[4:8] little-endian int32，存入 0.01 degree 多圈位置。"),
            ("move_motor_relative", "delta_degrees、max_speed_dps；可選延遲欄位", "bool", "首次先讀位置；將 delta 乘 100 後累加至 position_target_cdeg，再以 0xA4 送出目標位置。小於 0.005 degree 會忽略。"),
            ("stop_motor_immediately", "無", "bool", "送 0x81 並等待 ACK/回覆，同時清除漸停狀態與速度記錄。"),
            ("begin_soft_stop", "無", "無", "速度模式非零時啟用漸停；已經靜止時直接呼叫立即停止。"),
            ("update_soft_stop", "無", "無", "每 100 ms 將速度朝 0 改變 10 dps；到 0 後發 0x81。loop 會持續呼叫它。"),
            ("valid_speed", "speed_dps: int", "bool", "檢查 m 命令速度是否在 -720 至 720 dps。"),
            ("handle_command", "line: char*", "無", "剖析 r、m、n、s；n 可接一般 3 欄格式或含 sequence/時間戳的 6 欄延遲格式。"),
            ("setup", "無", "無", "初始化 USB Serial、印出指令說明，並呼叫 start_can。"),
            ("loop", "無", "無", "執行漸停、逐字讀 USB Serial、以換行或 100 ms 無新字元作為命令結束，再交給 handle_command。"),
        ],
        [1.42, 1.52, 1.02, 3.44],
        font_size=7.95,
    )
    add_table(
        doc,
        ["CAN 命令", "8-byte 資料格式", "程式如何編碼"],
        [
            ("0x9A 讀狀態", "[9A, 00, 00, 00, 00, 00, 00, 00]", "讀回溫度、0.01 V 電壓、錯誤碼。"),
            ("0x92 讀多圈角度", "[92, 00, 00, 00, 00, 00, 00, 00]", "讀回 bytes 4 至 7 的 little-endian int32，單位 0.01 degree。"),
            ("0xA2 速度控制", "[A2, 00, 00, 00, v0, v1, v2, v3]", "speed_dps × 100 轉 int32，v0 是最低位元；負值代表反向。"),
            ("0xA4 位置控制", "[A4, 00, s0, s1, p0, p1, p2, p3]", "s0:s1 為最大速度 dps 的 little-endian uint16；p0:p3 為多圈目標位置 cdeg。"),
            ("0x81 停止", "[81, 00, 00, 00, 00, 00, 00, 00]", "立即停止命令；s 在速度模式會先執行軟性減速。"),
        ],
        [1.55, 2.9, 3.05],
        font_size=8.35,
    )

    add_heading(doc, "9. 建議操作流程")
    add_heading(doc, "9.1 啟動前檢查", level=2)
    add_numbered(
        doc,
        [
            "確認 Xsens 電腦與控制電腦網路互通，且遠端 SSH 可連線。",
            "確認 ESP32、WCMCU230、馬達電源與 CANH/CANL 接線完成。",
            "在斷電狀態量測 CANH 與 CANL，確認等效終端電阻約為 60 ohm。",
            "關閉 Arduino Serial Monitor、VS Code Serial Monitor 等會占用 COM7 的程式。",
            "先以 ESP32 的 r 指令確認能收到馬達狀態回覆，再啟用自動跟隨。",
        ],
    )
    add_heading(doc, "9.2 一般啟動", level=2)
    add_code(doc, "cd C:\\Users\\e11512\\Documents\\Codex\\smr\\SMR-\npython python\\control_system.py --arm\n\nsystem> start")
    add_body(doc, "若只想確認 Xsens 資料、網路與控制決策，而不控制馬達，省略 --arm 即可進入 Dry Run。")
    add_heading(doc, "9.3 停止", level=2)
    add_code(doc, "system> end")
    add_body(doc, "也可使用 Ctrl+C 結束。正常結束時，bridge 與控制器都會嘗試送出 s。任何異常情況下仍應保留可立即切斷馬達外部電源的實體安全方式。")

    add_heading(doc, "10. 後續改善方向")
    add_bullets(
        doc,
        [
            "將 watchdog 改為資料中斷後主動發送停止命令，並在 CAN 無 ACK 時停止重試。",
            "利用馬達編碼器回授確認實際位置，建立 yaw 與馬達位置的閉迴路誤差補償。",
            "使用 NTP 或 PTP 同步兩台電腦時鐘，提高端到端延遲紀錄可信度。",
            "以設定檔統一 IP、COM 與控制限制，並建立 CAN 故障保護狀態。",
        ],
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = "SMR 程式運作框架"
    doc.core_properties.subject = "Xsens MTi-630R 與 ESP32 CAN 馬達控制系統"
    doc.core_properties.author = "SMR Team"
    doc.save(OUTPUT)
    architecture_diagram.unlink(missing_ok=True)
    print(OUTPUT.resolve())


if __name__ == "__main__":
    build_document()
